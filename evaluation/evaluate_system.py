"""
evaluation/evaluate_system.py

Measures Industrial Mind detection efficiency OFFLINE (separate from the
dashboard).

1. Reads the controlled test cases from evaluation/test_data.xlsx (TestCases).
2. Runs every case through the REAL middleware security code:
       docs.hmac_auth (frame build/parse)  ->  middleware.cyber_agent.check()
       middleware.physical_agent.check_command()  ->  middleware.orchestrator.decide()
   and writes the Cyber / Physical verdicts, final decision and
   Predicted Label back into the sheet.
3. Computes TP, TN, FP, FN, Accuracy, Precision, Recall, F1, Detection Rate,
   FPR, Specificity and the confusion matrix from Actual vs Predicted Label.
4. Rebuilds the "Results" and "Confusion Matrix" sheets. Every metric cell is
   an Excel formula over the TestCases sheet (so it recalculates if you edit
   labels in Excel); next to it is the value computed by this script, and a
   check column confirming both agree.
5. Saves evaluation/confusion_matrix.png and embeds it in the workbook.

Nothing is hard-coded: change or add rows in TestCases and rerun.

    python evaluation/evaluate_system.py              # rerun cases + metrics
    python evaluation/evaluate_system.py --no-rerun   # metrics from the existing
                                                      # Predicted Label column only
"""

import argparse
import json
import os
import secrets
import sys
import time
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from openpyxl import load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

from docs import attack_log, hmac_auth  # noqa: E402
from docs.interfaces import MACHINES, MACHINE_BY_NAME, make_command, to_raw  # noqa: E402
from middleware import cyber_agent, orchestrator, physical_agent  # noqa: E402

DEFAULT_BOOK = os.path.join(HERE, "test_data.xlsx")
DEFAULT_PNG = os.path.join(HERE, "confusion_matrix.png")

POSITIVE, NEGATIVE = "ATTACK", "NORMAL"

OUTPUT_COLS = [
    "Cyber Verdict", "Cyber Reason", "Physical Verdict", "Physical Reason",
    "Final Decision", "Orchestrator Reason", "Detected Attack Type",
    "Predicted Label", "Outcome",
]

ARIAL = Font(name="Arial")
ARIAL_B = Font(name="Arial", bold=True)
HDR = Font(name="Arial", bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="385723")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


# ---------------------------------------------------------------------------
# Running one case through the real pipeline
# ---------------------------------------------------------------------------

def _resolve_register(value):
    s = str(value).strip()
    if s in MACHINE_BY_NAME:
        return MACHINE_BY_NAME[s]["register"]
    return int(float(s))


def _num(v, default=0.0):
    if v is None or v == "":
        return default
    return float(v)


def _signed_command(tx_id, ts, ip, fc, register, value, mode):
    """Builds the command exactly as the middleware would after parsing the wire frame."""
    now_ms = int(ts * 1000)

    if mode == "missing":
        return make_command(tx_id, ts, ip, fc, register, value, raw_value=to_raw(value), auth=None)

    key = secrets.token_hex(32).encode() if mode == "wrong_key" else None
    ts_ms = now_ms - 300_000 if mode == "expired" else now_ms

    frame = hmac_auth.build_frame(register, value, key=key, timestamp_ms=ts_ms)

    if mode == "tampered":
        frame[2] = min(65535, frame[2] + 790)   # +79 units after signing

    parsed = hmac_auth.parse_frame(frame)

    if mode == "malformed":
        parsed["auth"]["mac"] = parsed["auth"]["mac"][:10]

    raw = parsed["raw_value"]
    return make_command(tx_id, ts, ip, fc, register, raw / 10.0, raw_value=raw, auth=parsed["auth"])


def run_case(row):
    cyber_agent.reset_state()

    register = _resolve_register(row["Machine Register"])
    value = _num(row["Requested Value"])
    ip = str(row["Source IP"]).strip()
    fc = int(_num(row["Function Code"], 16))
    mode = str(row["Signature Mode"] or "valid").strip().lower()
    repeat = max(1, int(_num(row["Repeat Count"], 1)))
    dt = _num(row["Seconds Since Last Settle"], 12.0)

    registers = {
        m["register"]: _num(row.get(f"Current {m['name']}"), m["initial"]) for m in MACHINES
    }
    setpoints = dict(registers)
    pending = str(row.get("Pending Setpoints (JSON)") or "").strip()
    if pending:
        for name, v in json.loads(pending).items():
            setpoints[_resolve_register(name)] = float(v)

    now = time.time()
    state = {"registers": registers, "setpoints": setpoints, "timestamp": now - dt}

    # Build the send sequence. Earlier sends in a burst differ slightly so
    # the evaluated (last) command is exactly the requested value.
    sends = []
    if mode == "replayed":
        cmd = _signed_command(f"{row['Case ID']}-0", now, ip, fc, register, value, "valid")
        sends = [cmd, dict(cmd, tx_id=f"{row['Case ID']}-1", timestamp=now + 0.05)]
    else:
        for i in range(repeat):
            v = round(value - 0.1 * (repeat - 1 - i), 1)
            sends.append(_signed_command(f"{row['Case ID']}-{i}", now + i * 0.01, ip, fc, register, v, mode))

    for cmd in sends:
        start = time.time()
        cyber = cyber_agent.check(cmd)
        physical = physical_agent.check_command(cmd, state)
        decision = orchestrator.decide(cyber, physical, start)

    predicted = POSITIVE if decision["verdict"] == "DROP" else NEGATIVE

    return {
        "Cyber Verdict": "PASS" if cyber["pass"] else "FAIL",
        "Cyber Reason": cyber["reason"],
        "Physical Verdict": "PASS" if physical["pass"] else "FAIL",
        "Physical Reason": physical["reason"],
        "Final Decision": decision["verdict"],
        "Orchestrator Reason": decision["reason"],
        "Detected Attack Type": attack_log.classify_reason(decision["reason"]) if predicted == POSITIVE else "-",
        "Predicted Label": predicted,
    }


# ---------------------------------------------------------------------------
# Metrics
# ---------------------------------------------------------------------------

def compute_metrics(actual, predicted):
    pairs = [(a, p) for a, p in zip(actual, predicted) if a in (POSITIVE, NEGATIVE) and p in (POSITIVE, NEGATIVE)]
    tp = sum(1 for a, p in pairs if a == POSITIVE and p == POSITIVE)
    tn = sum(1 for a, p in pairs if a == NEGATIVE and p == NEGATIVE)
    fp = sum(1 for a, p in pairs if a == NEGATIVE and p == POSITIVE)
    fn = sum(1 for a, p in pairs if a == POSITIVE and p == NEGATIVE)
    total = tp + tn + fp + fn

    def div(a, b):
        return a / b if b else 0.0

    precision = div(tp, tp + fp)
    recall = div(tp, tp + fn)
    return {
        "Total Cases": total,
        "Actual Attacks": tp + fn,
        "Actual Normal": tn + fp,
        "TP": tp, "TN": tn, "FP": fp, "FN": fn,
        "Accuracy": div(tp + tn, total),
        "Precision": precision,
        "Recall": recall,
        "F1-score": div(2 * precision * recall, precision + recall),
        "Detection Rate": recall,
        "False Positive Rate": div(fp, fp + tn),
        "Specificity": div(tn, tn + fp),
    }


# ---------------------------------------------------------------------------
# Workbook output
# ---------------------------------------------------------------------------

def _header_map(ws):
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def _style_header(cell):
    cell.font = HDR
    cell.fill = HDR_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BOX


def write_results(wb, cols, last_row, metrics, attack_types):
    for name in ("Results", "Confusion Matrix"):
        if name in wb.sheetnames:
            del wb[name]

    L = {h: get_column_letter(c) for h, c in cols.items()}
    act = f"TestCases!${L['Actual Label']}$2:${L['Actual Label']}${last_row}"
    pred = f"TestCases!${L['Predicted Label']}$2:${L['Predicted Label']}${last_row}"
    exp = f"TestCases!${L['Expected Attack Type']}$2:${L['Expected Attack Type']}${last_row}"

    ws = wb.create_sheet("Results", 1)
    ws["A1"] = "Industrial Mind - System Efficiency Results"
    ws["A1"].font = Font(name="Arial", bold=True, size=14)
    ws["A2"] = (f"Evaluated {datetime.now().isoformat(timespec='seconds')} over TestCases rows 2-{last_row}. "
                "Positive class = ATTACK. Formulas recalculate from TestCases.")
    ws["A2"].font = Font(name="Arial", italic=True, color="595959")

    for i, h in enumerate(["Metric", "Value (Excel formula)", "Value (evaluate_system.py)", "Check", "Definition"], 1):
        _style_header(ws.cell(row=4, column=i, value=h))

    rows = [
        ("Total Cases", f'=COUNTIF({act},"ATTACK")+COUNTIF({act},"NORMAL")', "0", "labelled cases"),
        ("Actual Attacks", f'=COUNTIF({act},"ATTACK")', "0", "Actual Label = ATTACK"),
        ("Actual Normal", f'=COUNTIF({act},"NORMAL")', "0", "Actual Label = NORMAL"),
        ("TP", f'=COUNTIFS({act},"ATTACK",{pred},"ATTACK")', "0", "attack correctly dropped"),
        ("TN", f'=COUNTIFS({act},"NORMAL",{pred},"NORMAL")', "0", "normal command correctly passed"),
        ("FP", f'=COUNTIFS({act},"NORMAL",{pred},"ATTACK")', "0", "normal command wrongly dropped"),
        ("FN", f'=COUNTIFS({act},"ATTACK",{pred},"NORMAL")', "0", "attack wrongly passed"),
        ("Accuracy", "=IFERROR((B8+B9)/(B8+B9+B10+B11),0)", "0.0%", "(TP+TN)/(TP+TN+FP+FN)"),
        ("Precision", "=IFERROR(B8/(B8+B10),0)", "0.0%", "TP/(TP+FP)"),
        ("Recall", "=IFERROR(B8/(B8+B11),0)", "0.0%", "TP/(TP+FN)"),
        ("F1-score", "=IFERROR(2*B13*B14/(B13+B14),0)", "0.0%", "2*Precision*Recall/(Precision+Recall)"),
        ("Detection Rate", "=IFERROR(B8/(B8+B11),0)", "0.0%", "attacks detected / all attacks (= Recall)"),
        ("False Positive Rate", "=IFERROR(B10/(B10+B9),0)", "0.0%", "FP/(FP+TN)"),
        ("Specificity", "=IFERROR(B9/(B9+B10),0)", "0.0%", "TN/(TN+FP)"),
    ]
    for r, (name, formula, fmt, definition) in enumerate(rows, start=5):
        ws.cell(row=r, column=1, value=name).font = ARIAL_B
        b = ws.cell(row=r, column=2, value=formula)
        c = ws.cell(row=r, column=3, value=metrics[name])
        b.number_format = c.number_format = fmt
        b.font = c.font = ARIAL
        ws.cell(row=r, column=4, value=f'=IF(ABS(B{r}-C{r})<0.000001,"OK","MISMATCH")').font = ARIAL
        ws.cell(row=r, column=5, value=definition).font = Font(name="Arial", color="595959")
        for col in range(1, 5):
            ws.cell(row=r, column=col).border = BOX

    start = 5 + len(rows) + 2
    ws.cell(row=start - 1, column=1, value="Detection by attack type").font = Font(name="Arial", bold=True, size=12)
    for i, h in enumerate(["Expected Attack Type", "Attack Cases", "Detected (DROP)", "Detection Rate"], 1):
        _style_header(ws.cell(row=start, column=i, value=h))
    for k, t in enumerate(attack_types, start=start + 1):
        ws.cell(row=k, column=1, value=t).font = ARIAL
        ws.cell(row=k, column=2, value=f'=COUNTIFS({exp},A{k},{act},"ATTACK")').font = ARIAL
        ws.cell(row=k, column=3, value=f'=COUNTIFS({exp},A{k},{act},"ATTACK",{pred},"ATTACK")').font = ARIAL
        d = ws.cell(row=k, column=4, value=f"=IFERROR(C{k}/B{k},0)")
        d.font, d.number_format = ARIAL, "0.0%"
        for col in range(1, 5):
            ws.cell(row=k, column=col).border = BOX

    for letter, w in zip("ABCDE", (36, 22, 26, 12, 48)):
        ws.column_dimensions[letter].width = w

    # ---- Confusion matrix ----
    cm = wb.create_sheet("Confusion Matrix", 2)
    cm["A1"] = "Confusion Matrix (positive class = ATTACK)"
    cm["A1"].font = Font(name="Arial", bold=True, size=14)
    _style_header(cm.cell(row=3, column=1, value="Actual \\ Predicted"))
    _style_header(cm.cell(row=3, column=2, value="Predicted ATTACK"))
    _style_header(cm.cell(row=3, column=3, value="Predicted NORMAL"))
    _style_header(cm.cell(row=3, column=4, value="Total"))
    grid = [
        ("Actual ATTACK", "=Results!B8", "=Results!B11"),
        ("Actual NORMAL", "=Results!B10", "=Results!B9"),
    ]
    for r, (lab, f1, f2) in enumerate(grid, start=4):
        _style_header(cm.cell(row=r, column=1, value=lab))
        for c, f in ((2, f1), (3, f2)):
            cell = cm.cell(row=r, column=c, value=f)
            cell.font = Font(name="Arial", bold=True, size=14)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BOX
        t = cm.cell(row=r, column=4, value=f"=B{r}+C{r}")
        t.font, t.alignment, t.border = ARIAL, Alignment(horizontal="center"), BOX
    _style_header(cm.cell(row=6, column=1, value="Total"))
    for c in (2, 3, 4):
        L2 = get_column_letter(c)
        cell = cm.cell(row=6, column=c, value=f"={L2}4+{L2}5")
        cell.font, cell.alignment, cell.border = ARIAL, Alignment(horizontal="center"), BOX
    cm["A8"] = "TP = Actual ATTACK & Predicted ATTACK · FN = Actual ATTACK & Predicted NORMAL · FP = Actual NORMAL & Predicted ATTACK · TN = Actual NORMAL & Predicted NORMAL"
    cm["A8"].font = Font(name="Arial", color="595959")
    for letter, w in zip("ABCD", (22, 20, 20, 12)):
        cm.column_dimensions[letter].width = w
    for r in (4, 5):
        cm.row_dimensions[r].height = 36

    return cm


def save_png(metrics, path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.colors import LinearSegmentedColormap

    # Sequential single-hue blue ramp (light -> dark).
    cmap = LinearSegmentedColormap.from_list("im_blue", ["#cde2fb", "#86b6ef", "#3987e5", "#1c5cab", "#0d366b"])
    m = [[metrics["TP"], metrics["FN"]], [metrics["FP"], metrics["TN"]]]
    names = [["TP", "FN"], ["FP", "TN"]]
    vmax = max(1, max(max(r) for r in m))

    fig, ax = plt.subplots(figsize=(5.4, 4.6), dpi=150)
    ax.imshow(m, cmap=cmap, vmin=0, vmax=vmax)
    for i in range(2):
        for j in range(2):
            dark = m[i][j] / vmax > 0.55
            ax.text(j, i, f"{m[i][j]}\n{names[i][j]}", ha="center", va="center",
                    fontsize=15, fontweight="bold", color="#ffffff" if dark else "#1f2937")
    ax.set_xticks([0, 1], ["ATTACK", "NORMAL"])
    ax.set_yticks([0, 1], ["ATTACK", "NORMAL"])
    ax.set_xlabel("Predicted label", color="#374151")
    ax.set_ylabel("Actual label", color="#374151")
    ax.set_title(
        f"Industrial Mind confusion matrix (n={metrics['Total Cases']})\n"
        f"Accuracy {metrics['Accuracy']:.1%} · Precision {metrics['Precision']:.1%} · "
        f"Recall {metrics['Recall']:.1%} · F1 {metrics['F1-score']:.1%}",
        fontsize=9.5, color="#111827",
    )
    for s in ax.spines.values():
        s.set_visible(False)
    ax.tick_params(length=0, colors="#374151")
    fig.tight_layout()
    fig.savefig(path)
    plt.close(fig)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description="Evaluate Industrial Mind detection efficiency")
    ap.add_argument("--book", default=DEFAULT_BOOK)
    ap.add_argument("--png", default=DEFAULT_PNG)
    ap.add_argument("--no-rerun", action="store_true",
                    help="do not rerun the agents; use the Predicted Label column as-is")
    ap.add_argument("--no-png", action="store_true")
    args = ap.parse_args()

    # Offline run: never push evaluation traffic to a live dashboard.
    orchestrator._forward_decision_to_dashboard = lambda payload: None

    wb = load_workbook(args.book)
    ws = wb["TestCases"]
    cols = _header_map(ws)

    for h in OUTPUT_COLS:
        if h not in cols:
            c = ws.max_column + 1
            ws.cell(row=1, column=c, value=h)
            _style_header(ws.cell(row=1, column=c))
            cols[h] = c

    data_rows = [r for r in range(2, ws.max_row + 1)
                 if ws.cell(row=r, column=cols["Case ID"]).value not in (None, "")]
    if not data_rows:
        sys.exit("No test cases found in TestCases.")
    last_row = data_rows[-1]

    L = {h: get_column_letter(c) for h, c in cols.items()}
    actual, predicted = [], []

    print(f"{'Case':6} {'Actual':7} {'Pred':7} {'Out':3}  Decision reason")
    for r in data_rows:
        row = {h: ws.cell(row=r, column=c).value for h, c in cols.items()}
        a = str(row["Actual Label"] or "").strip().upper()

        if args.no_rerun:
            p = str(row.get("Predicted Label") or "").strip().upper()
            reason = row.get("Orchestrator Reason") or ""
        else:
            try:
                result = run_case(row)
            except Exception as e:   # bad input row -> visible, not silently skipped
                result = {h: "" for h in OUTPUT_COLS}
                result["Orchestrator Reason"] = f"CASE ERROR: {e}"
            for h, v in result.items():
                cell = ws.cell(row=r, column=cols[h], value=v)
                cell.font = ARIAL
            p = result.get("Predicted Label", "")
            reason = result.get("Orchestrator Reason", "")

        # Outcome stays a live formula over Actual / Predicted.
        ac, pc = f"{L['Actual Label']}{r}", f"{L['Predicted Label']}{r}"
        ws.cell(row=r, column=cols["Outcome"], value=(
            f'=IF(OR({pc}="",{ac}=""),"",IF({ac}="ATTACK",IF({pc}="ATTACK","TP","FN"),'
            f'IF({pc}="ATTACK","FP","TN")))'
        )).font = ARIAL

        actual.append(a)
        predicted.append(p)
        out = {("ATTACK", "ATTACK"): "TP", ("NORMAL", "NORMAL"): "TN",
               ("NORMAL", "ATTACK"): "FP", ("ATTACK", "NORMAL"): "FN"}.get((a, p), "?")
        print(f"{row['Case ID']:6} {a:7} {p:7} {out:3}  {str(reason)[:90]}")

    metrics = compute_metrics(actual, predicted)

    attack_types = sorted({
        str(ws.cell(row=r, column=cols["Expected Attack Type"]).value).strip()
        for r, a in zip(data_rows, actual)
        if a == POSITIVE and ws.cell(row=r, column=cols["Expected Attack Type"]).value not in (None, "", "-")
    })

    cm = write_results(wb, cols, last_row, metrics, attack_types)

    if not args.no_png:
        save_png(metrics, args.png)
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(args.png)
            img.width, img.height = 486, 414
            cm.add_image(img, "F3")
        except Exception as e:
            print(f"(could not embed PNG in workbook: {e})")

    wb.calculation.fullCalcOnLoad = True
    wb.save(args.book)

    print("\n" + "=" * 46)
    for k in ("Total Cases", "TP", "TN", "FP", "FN"):
        print(f"{k:22} {metrics[k]}")
    for k in ("Accuracy", "Precision", "Recall", "F1-score", "Detection Rate", "False Positive Rate", "Specificity"):
        print(f"{k:22} {metrics[k]:.2%}")
    print("=" * 46)
    print(f"Confusion matrix      [[TP={metrics['TP']}, FN={metrics['FN']}], [FP={metrics['FP']}, TN={metrics['TN']}]]")
    print(f"Workbook updated      {args.book}")
    if not args.no_png:
        print(f"Confusion matrix PNG  {args.png}")


if __name__ == "__main__":
    main()
