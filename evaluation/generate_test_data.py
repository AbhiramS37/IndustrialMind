"""
evaluation/generate_test_data.py

Creates evaluation/test_data.xlsx with controlled NORMAL and ATTACK test
cases (inputs + ground-truth Actual Label only).

The Predicted Label, agent verdicts and final decision columns are left empty
here; evaluation/evaluate_system.py fills them by running every case through
the REAL Cyber Agent, Physical Agent and Orchestrator.

    python evaluation/generate_test_data.py          # refuses to overwrite
    python evaluation/generate_test_data.py --force  # regenerate from scratch

You can also edit/add rows directly in the workbook and just rerun
evaluate_system.py.
"""

import argparse
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from docs.interfaces import MACHINES  # noqa: E402

OUT = os.path.join(HERE, "test_data.xlsx")

STATE_COLS = [f"Current {m['name']}" for m in MACHINES]

INPUT_COLS = [
    "Case ID", "Scenario", "Actual Label", "Expected Attack Type",
    "Source IP", "Function Code", "Machine Register", "Requested Value",
    "Signature Mode", "Repeat Count", "Seconds Since Last Settle",
] + STATE_COLS + ["Pending Setpoints (JSON)", "Description"]

OUTPUT_COLS = [
    "Cyber Verdict", "Cyber Reason", "Physical Verdict", "Physical Reason",
    "Final Decision", "Orchestrator Reason", "Detected Attack Type",
    "Predicted Label", "Outcome",
]

SIGNATURE_MODES = ["valid", "missing", "tampered", "wrong_key", "expired", "replayed", "malformed"]

DEFAULT_STATE = {m["name"]: m["initial"] for m in MACHINES}


def case(cid, scenario, label, attack_type, register, value, *, ip="127.0.0.1", fc=16,
         sig="valid", repeat=1, dt=12.0, state=None, pending=None, desc=""):
    st = dict(DEFAULT_STATE)
    st.update(state or {})
    return {
        "Case ID": cid, "Scenario": scenario, "Actual Label": label,
        "Expected Attack Type": attack_type, "Source IP": ip, "Function Code": fc,
        "Machine Register": register, "Requested Value": value,
        "Signature Mode": sig, "Repeat Count": repeat, "Seconds Since Last Settle": dt,
        **{f"Current {k}": v for k, v in st.items()},
        "Pending Setpoints (JSON)": json.dumps(pending) if pending else "",
        "Description": desc,
    }


N, A = "NORMAL", "ATTACK"

CASES = [
    # ---------------- NORMAL operation ----------------
    case("N01", "Operator pressure step", N, "-", "TANK_PRESSURE", 51.0, desc="50 -> 51 PSI signed operator step"),
    case("N02", "Operator conveyor step", N, "-", "CONVEYOR_SPEED", 2.0, desc="0 -> 2 RPM"),
    case("N03", "Operator cooling step", N, "-", "COOLING_VALVE", 46.0, desc="45 -> 46 deg"),
    case("N04", "Operator heater step", N, "-", "HEATER_TEMP", 61.0, desc="60 -> 61 C"),
    case("N05", "Operator feed pump step", N, "-", "FEED_PUMP", 41.0, desc="40 -> 41 %"),
    case("N06", "Operator relief valve step", N, "-", "RELIEF_VALVE", 46.0, desc="45 -> 46 deg"),
    case("N07", "Heater high with cooling open", N, "-", "HEATER_TEMP", 112.0, dt=10,
         state={"HEATER_TEMP": 102.0, "COOLING_VALVE": 60.0},
         desc="Legitimate multi-register combination: hot reactor but adequate cooling"),
    case("N08", "Pump high with relief open", N, "-", "FEED_PUMP", 75.0, dt=10,
         state={"FEED_PUMP": 65.0, "RELIEF_VALVE": 50.0},
         desc="Legitimate multi-register combination: high feed with relief path open"),
    case("N09", "Relief nearly closed at low feed", N, "-", "RELIEF_VALVE", 12.0, dt=10,
         state={"RELIEF_VALVE": 45.0, "FEED_PUMP": 40.0},
         desc="Relief low is safe because feed pump is low"),
    case("N10", "Cooling low at moderate heat", N, "-", "COOLING_VALVE", 20.0, dt=10,
         state={"COOLING_VALVE": 30.0, "HEATER_TEMP": 80.0},
         desc="Cooling low is safe because heater setpoint is moderate"),
    case("N11", "Pressure up with warm heater", N, "-", "TANK_PRESSURE", 86.0, dt=5,
         state={"TANK_PRESSURE": 80.0, "HEATER_TEMP": 110.0},
         desc="Below HOT_PRESSURIZED_VESSEL heater threshold (125 C)"),
    case("N12", "Corrective move out of unsafe state", N, "-", "FEED_PUMP", 75.0, dt=10,
         state={"FEED_PUMP": 80.0, "RELIEF_VALVE": 10.0},
         desc="Plant already in OVERPRESSURE combination; operator reduces pump"),
    case("N13", "Large move after long idle", N, "-", "CONVEYOR_SPEED", 80.0, dt=10,
         desc="0 -> 80 RPM, plant had 10 s to reach it (rate 10 RPM/s)"),
    case("N14", "Second trusted SCADA host", N, "-", "COOLING_VALVE", 46.0, ip="192.168.1.10",
         desc="Signed command from the other whitelisted SCADA address"),
    case("N15", "Pressure and heat both high but below limits", N, "-", "HEATER_TEMP", 124.0, dt=10,
         state={"HEATER_TEMP": 115.0, "TANK_PRESSURE": 84.0, "COOLING_VALVE": 60.0}),
    case("N16", "Fast conveyor, moderate pressure", N, "-", "CONVEYOR_SPEED", 110.0, dt=5,
         state={"CONVEYOR_SPEED": 100.0, "TANK_PRESSURE": 80.0}),
    case("N17", "Cooling reduced at normal pressure", N, "-", "COOLING_VALVE", 44.0, dt=5,
         state={"COOLING_VALVE": 60.0, "TANK_PRESSURE": 30.0}),
    case("N18", "Emergency fast cooling open", N, "-", "COOLING_VALVE", 80.0, dt=1,
         desc="Legitimate emergency action faster than the valve ramp (edge case)"),
    case("N19", "Short operator burst", N, "-", "RELIEF_VALVE", 46.0, repeat=3,
         desc="3 quick legitimate adjustments (below 5 req/s)"),
    case("N20", "Operator relief step", N, "-", "RELIEF_VALVE", 51.0, state={"RELIEF_VALVE": 50.0}),

    # ---------------- ATTACKS ----------------
    case("A01", "Unsigned command", A, "Missing HMAC (Unsigned Command)", "COOLING_VALVE", 46.0,
         fc=6, sig="missing", desc="Plain FC6 write with no HMAC"),
    case("A02", "Value tampered in transit", A, "Tampered / Forged HMAC", "HEATER_TEMP", 61.0,
         sig="tampered", desc="Signed 61 C, value changed to 140 C, original tag kept"),
    case("A03", "Forged signature", A, "Tampered / Forged HMAC", "COOLING_VALVE", 46.0,
         sig="wrong_key", desc="Signed with an attacker-chosen key"),
    case("A04", "Delayed replay (expired)", A, "Expired Signature (Delayed Replay)", "CONVEYOR_SPEED", 2.0,
         sig="expired", desc="Correct signature, timestamp 5 minutes old"),
    case("A05", "Immediate replay", A, "Replay", "CONVEYOR_SPEED", 2.0,
         sig="replayed", desc="Byte-identical signed frame sent twice; second is evaluated"),
    case("A06", "Malformed signed frame", A, "Malformed Signed Frame", "COOLING_VALVE", 46.0,
         sig="malformed", desc="Truncated HMAC tag"),
    case("A07", "Rogue host with stolen key", A, "Untrusted Source", "COOLING_VALVE", 46.0,
         ip="10.0.0.99", desc="Valid signature but connection from a non-SCADA address"),
    case("A08", "Probe of unmapped register", A, "Invalid Register Access", "7", 10.0,
         desc="Write to address 7 (no such machine)"),
    case("A09", "Invalid function code", A, "Invalid Function Code", "CONVEYOR_SPEED", 2.0, fc=5,
         desc="FC5 (write coil) against a holding register"),
    case("A10", "Command flood", A, "Flood / Rate Limit", "CONVEYOR_SPEED", 2.0, repeat=8,
         desc="8 signed commands in < 1 s; last one evaluated"),
    case("A11", "Impossible tank pressure", A, "Impossible Command (Out of Bounds)", "TANK_PRESSURE", 999.0),
    case("A12", "Impossible heater temperature", A, "Impossible Command (Out of Bounds)", "HEATER_TEMP", 180.0),
    case("A13", "Physically impossible ramp", A, "Rate-of-Change Violation", "CONVEYOR_SPEED", 100.0, dt=0.2,
         desc="0 -> 100 RPM in 0.2 s"),
    case("A14", "Thermal runaway via heater", A, "Cross-Register Correlation", "HEATER_TEMP", 112.0, dt=10,
         state={"HEATER_TEMP": 102.0, "COOLING_VALVE": 20.0},
         desc="Each value safe alone; heater 112 + cooling 20 is unsafe"),
    case("A15", "Thermal runaway via cooling valve", A, "Cross-Register Correlation", "COOLING_VALVE", 24.0, dt=10,
         state={"HEATER_TEMP": 115.0, "COOLING_VALVE": 30.0}),
    case("A16", "Overpressure via feed pump", A, "Cross-Register Correlation", "FEED_PUMP", 75.0, dt=10,
         state={"FEED_PUMP": 65.0, "RELIEF_VALVE": 10.0}),
    case("A17", "Overpressure via pending setpoint", A, "Cross-Register Correlation", "FEED_PUMP", 75.0, dt=10,
         state={"FEED_PUMP": 65.0, "RELIEF_VALVE": 40.0}, pending={"RELIEF_VALVE": 10.0},
         desc="Relief telemetry still 40 but an approved setpoint of 10 is ramping"),
    case("A18", "Hot pressurized vessel", A, "Cross-Register Correlation", "TANK_PRESSURE", 86.0, dt=5,
         state={"TANK_PRESSURE": 80.0, "HEATER_TEMP": 130.0}),
    case("A19", "Valve slam at low pressure", A, "Cross-Register Correlation", "COOLING_VALVE", 30.0, dt=10,
         state={"TANK_PRESSURE": 8.0, "COOLING_VALVE": 60.0}),
    case("A20", "Conveyor max at pressure limit", A, "Cross-Register Correlation", "CONVEYOR_SPEED", 110.0, dt=5,
         state={"CONVEYOR_SPEED": 100.0, "TANK_PRESSURE": 95.0}),
    case("A21", "Insider setpoint drift", A, "Stealthy In-Bounds Manipulation", "CONVEYOR_SPEED", 90.0, dt=5,
         state={"CONVEYOR_SPEED": 80.0},
         desc="Compromised HMI nudges conveyor within all limits (edge case)"),
    case("A22", "First step of low-and-slow correlation attack", A, "Cross-Register Correlation",
         "HEATER_TEMP", 70.0, dt=10, state={"HEATER_TEMP": 60.0, "COOLING_VALVE": 20.0},
         desc="Individually safe preparatory step of the A14 attack (edge case)"),
]


def build(path):
    wb = Workbook()

    ws = wb.active
    ws.title = "TestCases"
    headers = INPUT_COLS + OUTPUT_COLS
    ws.append(headers)

    hdr_font = Font(name="Arial", bold=True, color="FFFFFF")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font = hdr_font
        c.fill = PatternFill("solid", fgColor="1F4E78" if h in INPUT_COLS else "385723")
        c.alignment = Alignment(wrap_text=True, vertical="center")

    input_font = Font(name="Arial", color="0000FF")
    for r in CASES:
        ws.append([r.get(h, "") for h in INPUT_COLS] + [""] * len(OUTPUT_COLS))
        row = ws.max_row
        for i in range(1, len(INPUT_COLS) + 1):
            ws.cell(row=row, column=i).font = input_font
        for i in range(len(INPUT_COLS) + 1, len(headers) + 1):
            ws.cell(row=row, column=i).font = Font(name="Arial")

    last = ws.max_row + 200
    col = {h: ws.cell(row=1, column=i).column_letter for i, h in enumerate(headers, start=1)}
    dv_label = DataValidation(type="list", formula1='"NORMAL,ATTACK"', allow_blank=False)
    dv_sig = DataValidation(type="list", formula1='"' + ",".join(SIGNATURE_MODES) + '"')
    ws.add_data_validation(dv_label)
    ws.add_data_validation(dv_sig)
    dv_label.add(f"{col['Actual Label']}2:{col['Actual Label']}{last}")
    dv_sig.add(f"{col['Signature Mode']}2:{col['Signature Mode']}{last}")

    widths = {"Case ID": 8, "Scenario": 34, "Description": 55, "Cyber Reason": 45,
              "Physical Reason": 55, "Orchestrator Reason": 55, "Expected Attack Type": 32,
              "Detected Attack Type": 32, "Pending Setpoints (JSON)": 24}
    for h, letter in col.items():
        ws.column_dimensions[letter].width = widths.get(h, 15)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "C2"

    # ---- README / legend ----
    rd = wb.create_sheet("README")
    lines = [
        ("Industrial Mind - System Efficiency Test Data", True),
        ("", False),
        ("Blue columns (TestCases A..%s) are INPUTS you may edit or extend with new rows." % col["Description"], False),
        ("Green columns are OUTPUTS written by evaluation/evaluate_system.py - do not edit by hand.", False),
        ("Actual Label is the ground truth: NORMAL or ATTACK. Positive class = ATTACK.", False),
        ("Predicted Label = ATTACK when the real Orchestrator decision is DROP, NORMAL when PASS.", False),
        ("", False),
        ("Column meanings", True),
        ("Machine Register: machine name from docs/interfaces.MACHINES (e.g. HEATER_TEMP) or a raw address number.", False),
        ("Signature Mode: valid | missing | tampered | wrong_key | expired | replayed | malformed.", False),
        ("Repeat Count: number of back-to-back sends from the same source; the LAST one is evaluated (flood tests).", False),
        ("Seconds Since Last Settle: time the register has had to move since it last settled (rate-of-change check).", False),
        ("Current <MACHINE>: live PLC telemetry for each machine when the command arrives.", False),
        ('Pending Setpoints (JSON): approved targets still ramping, e.g. {"RELIEF_VALVE": 10}.', False),
        ("", False),
        ("Example row", True),
        ("A14 | heater 102 -> 112 C, cooling valve 20 deg, dt 10 s, valid signature | Actual ATTACK", False),
        ("", False),
        ("Rerun:  python evaluation/evaluate_system.py", True),
        ("Results and Confusion Matrix sheets are rebuilt with Excel formulas over TestCases on every run.", False),
    ]
    for text, bold in lines:
        rd.append([text])
        rd.cell(row=rd.max_row, column=1).font = Font(name="Arial", bold=bold, size=12 if bold else 10)
    rd.column_dimensions["A"].width = 120

    wb.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--force", action="store_true", help="overwrite an existing workbook")
    ap.add_argument("--out", default=OUT)
    args = ap.parse_args()

    if os.path.exists(args.out) and not args.force:
        print(f"{args.out} already exists - use --force to regenerate (your edits would be lost).")
        return

    build(args.out)
    print(f"Wrote {len(CASES)} test cases to {args.out}")
    print("Next: python evaluation/evaluate_system.py")


if __name__ == "__main__":
    main()
